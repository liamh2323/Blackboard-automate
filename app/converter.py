import os
import sys
import shutil
import tempfile
import subprocess
from pathlib import Path
from typing import Optional, Dict, Any, List

from bs4 import BeautifulSoup
import markdownify

from . import database as db
from .config import BASE_DIR, DOWNLOAD_DIR

# Ensure markitdown package from local repo is on sys.path
MARKITDOWN_SRC = BASE_DIR / "markitdown" / "packages" / "markitdown" / "src"
if MARKITDOWN_SRC.exists() and str(MARKITDOWN_SRC) not in sys.path:
    sys.path.insert(0, str(MARKITDOWN_SRC))

try:
    from markitdown import MarkItDown
    _md_converter = MarkItDown()
except Exception:
    _md_converter = None

SUPPORTED_CONVERTIBLE_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".pptx", ".ppt", ".xlsx", ".html", ".htm", ".txt", ".rtf"
}

def html_to_markdown(html_content: Any) -> str:
    """Converts raw HTML (such as announcements or description fields) to clean Markdown."""
    if not html_content:
        return ""
    if isinstance(html_content, dict):
        html_content = html_content.get("rawText") or html_content.get("text") or html_content.get("formattedText") or ""
    elif not isinstance(html_content, str):
        html_content = str(html_content)

    if not html_content or not html_content.strip():
        return ""
    try:
        # Pre-clean HTML with BeautifulSoup to remove scripts/styles
        soup = BeautifulSoup(html_content, "html.parser")
        for tag in soup(["script", "style"]):
            tag.decompose()
        cleaned_html = str(soup)
        md = markdownify.markdownify(cleaned_html, heading_style="ATX", autolinks=True)
        return md.strip()
    except Exception:
        # Fallback to plain text if markdownify encounters malformed tokens
        return BeautifulSoup(html_content, "html.parser").get_text().strip()

def _convert_docx_fallback(docx_path: Path) -> Optional[str]:
    """Fallback text extraction for Word .docx documents."""
    try:
        import docx
        doc = docx.Document(str(docx_path))
        lines = [f"# {docx_path.stem}\n"]
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                lines.append(text)
        for table in doc.tables:
            if not table.rows:
                continue
            lines.append("\n| " + " | ".join(cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells) + " |")
            lines.append("| " + " | ".join("---" for _ in table.rows[0].cells) + " |")
            for row in table.rows[1:]:
                lines.append("| " + " | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells) + " |")
            lines.append("\n")
        return "\n\n".join(lines) if len(lines) > 1 else None
    except Exception:
        return None

def _convert_pptx_fallback(pptx_path: Path) -> Optional[str]:
    """Fallback text extraction for PowerPoint .pptx presentations."""
    try:
        from pptx import Presentation
        prs = Presentation(str(pptx_path))
        slides_md = []
        for idx, slide in enumerate(prs.slides, 1):
            slide_lines = [f"## Slide {idx}"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_lines.append(f"- {text}")
            if len(slide_lines) > 1:
                slides_md.append("\n".join(slide_lines))
        if slides_md:
            return f"# {pptx_path.stem}\n\n" + "\n\n---\n\n".join(slides_md)
        return None
    except Exception:
        return None

def _convert_ppt_fallback(ppt_path: Path) -> Optional[str]:
    """Converts legacy binary .ppt files to PDF via LibreOffice then extracts text."""
    try:
        with tempfile.TemporaryDirectory() as tmp_dir:
            cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(ppt_path), "--outdir", tmp_dir]
            res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=25)
            if res.returncode == 0:
                tmp_pdf = Path(tmp_dir) / f"{ppt_path.stem}.pdf"
                if tmp_pdf.exists():
                    return _convert_pdf_fallback(tmp_pdf)
    except Exception:
        pass
    return None

def _convert_xlsx_fallback(xlsx_path: Path) -> Optional[str]:
    """Extracts tables from Excel .xlsx spreadsheets using openpyxl."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
        sections = [f"# {xlsx_path.stem}\n"]
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            sections.append(f"## Sheet: {sheet_name}\n")
            rows = list(sheet.iter_rows(values_only=True))
            non_empty_rows = [r for r in rows if any(c is not None for c in r)][:40]
            if non_empty_rows:
                header = [str(c if c is not None else "") for c in non_empty_rows[0]]
                sections.append("| " + " | ".join(header) + " |")
                sections.append("| " + " | ".join("---" for _ in header) + " |")
                for r in non_empty_rows[1:]:
                    sections.append("| " + " | ".join(str(c if c is not None else "").replace("\n", " ") for c in r) + " |")
                sections.append("\n")
        return "\n".join(sections)
    except Exception:
        return None

def _convert_pdf_fallback(pdf_path: Path) -> Optional[str]:
    """Fallback text extraction for PDF documents via pdfplumber."""
    try:
        import pdfplumber
        pages_md = []
        with pdfplumber.open(str(pdf_path)) as pdf:
            for idx, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    pages_md.append(f"<!-- Page {idx} -->\n{text.strip()}")
            if pages_md:
                return f"# {pdf_path.stem}\n\n" + "\n\n---\n\n".join(pages_md)
            elif len(pdf.pages) > 0:
                # Scanned handwritten or image document without text layer
                return f"# {pdf_path.stem}\n\n> *Note: Scanned handwritten or image document ({len(pdf.pages)} pages).*"
        return None
    except Exception:
        return None

def _convert_txt_fallback(txt_path: Path) -> Optional[str]:
    """Fallback text extraction for plain text files."""
    try:
        with open(txt_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return f"# {txt_path.stem}\n\n```text\n{content}\n```"
    except Exception:
        return None

def convert_document_to_markdown(file_path: Path, output_md_path: Optional[Path] = None) -> Optional[Path]:
    """
    Converts a document (PDF, DOCX, PPTX, PPT, XLSX, TXT, HTML) to a Markdown document saved side-by-side.
    Uses MarkItDown with format-specific fallbacks.
    """
    if not file_path.exists():
        return None

    ext = file_path.suffix.lower()
    if ext not in SUPPORTED_CONVERTIBLE_EXTENSIONS:
        return None

    if output_md_path is None:
        output_md_path = file_path.with_suffix(".md")

    # If markdown already exists and has content, avoid re-converting
    if output_md_path.exists() and output_md_path.stat().st_size > 0:
        return output_md_path

    markdown_text = None

    # Primary converter: MarkItDown (except for legacy .ppt which MarkItDown does not support)
    if _md_converter is not None and ext != ".ppt":
        try:
            res = _md_converter.convert(str(file_path))
            if res and res.text_content and res.text_content.strip():
                markdown_text = res.text_content
        except Exception:
            markdown_text = None

    # Fallbacks per extension
    if not markdown_text:
        if ext == ".pdf":
            markdown_text = _convert_pdf_fallback(file_path)
        elif ext in [".docx", ".doc"]:
            markdown_text = _convert_docx_fallback(file_path)
        elif ext == ".pptx":
            markdown_text = _convert_pptx_fallback(file_path)
        elif ext == ".ppt":
            markdown_text = _convert_ppt_fallback(file_path)
        elif ext == ".xlsx":
            markdown_text = _convert_xlsx_fallback(file_path)
        elif ext == ".txt":
            markdown_text = _convert_txt_fallback(file_path)

    if markdown_text and markdown_text.strip():
        try:
            output_md_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_md_path, "w", encoding="utf-8") as f:
                f.write(markdown_text.strip())
            return output_md_path
        except Exception:
            return None

    return None

# Alias for backwards compatibility
convert_pdf_to_markdown = convert_document_to_markdown

def convert_file_record(file_row: Dict[str, Any]) -> bool:
    """Converts a database file record if its extension is supported and not yet converted."""
    local_path_str = file_row.get("local_path")
    if not local_path_str:
        return False

    local_path = Path(local_path_str)
    if not local_path.exists():
        return False

    if local_path.suffix.lower() in SUPPORTED_CONVERTIBLE_EXTENSIONS:
        md_path = convert_document_to_markdown(local_path)
        if md_path:
            db.mark_file_converted(file_row["id"], str(md_path))
            return True

    return False

def convert_all_pending_files(scan_disk: bool = True) -> int:
    """
    Finds all unconverted files in the database AND on disk in DOWNLOAD_DIR,
    converting all applicable documents to Markdown.
    """
    converted_count = 0

    # 1. Convert pending records registered in database
    pending = db.get_unconverted_files(limit=500)
    for row in pending:
        if convert_file_record(row):
            converted_count += 1

    # 2. Convert any existing files directly in DOWNLOAD_DIR
    if scan_disk and DOWNLOAD_DIR.exists():
        for root, _, files in os.walk(DOWNLOAD_DIR):
            for f in sorted(files):
                if f.startswith(".") or f.endswith(".md"):
                    continue
                file_path = Path(root) / f
                if file_path.suffix.lower() in SUPPORTED_CONVERTIBLE_EXTENSIONS:
                    target_md = file_path.with_suffix(".md")
                    if not target_md.exists() or target_md.stat().st_size == 0:
                        res = convert_document_to_markdown(file_path, target_md)
                        if res:
                            converted_count += 1
                            # Update DB if record exists for this path
                            try:
                                with db.get_connection() as conn:
                                    conn.cursor().execute(
                                        "UPDATE files SET is_converted=1, markdown_path=? WHERE local_path=?",
                                        (str(res), str(file_path))
                                    )
                            except Exception:
                                pass

    return converted_count
